#!/usr/bin/env python3
"""生成 S-S07-C01 多来源竞品监控训练资料。

默认在脚本所在目录重建网页、JSON/CSV、说明文档、PDF 与两个 Excel 工作簿。
依赖：openpyxl、reportlab；生成扫描型 PDF 还需要 pdftoppm（Poppler）。
课堂成品不会由此脚本生成，本脚本只生成练习所需的原始资料与空白模板。
"""

from __future__ import annotations

import argparse
import csv
import json
import shutil
import subprocess
import tempfile
import textwrap
from copy import copy
from datetime import datetime
from pathlib import Path


CASE_ID = "S-S07-C01"
CAPTURE_BASE = "2026-09-03T09:00:00+09:00"
DEFAULT_BASE_URL = "https://carrycodedata.github.io/materials/sales/S07/S-S07-C01/00_原始资料"


def write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(textwrap.dedent(content).strip() + "\n", encoding="utf-8")


def build_pages(root: Path) -> None:
    static = root / "02_静态竞品官网"
    dynamic = root / "03_JS动态活动页"
    review = root / "04_评论平台训练页"

    static_css = """
    :root{--ink:#17221f;--green:#0d735d;--cream:#fff9ef;--line:#ded8ca}
    *{box-sizing:border-box}body{margin:0;font-family:-apple-system,BlinkMacSystemFont,"Noto Sans JP",sans-serif;color:var(--ink);background:var(--cream)}
    header{padding:20px 7%;display:flex;justify-content:space-between;align-items:center;background:#fff;border-bottom:1px solid var(--line)}
    .brand{font-size:22px;font-weight:750}.tag{color:var(--green);font-weight:700}.hero{padding:64px 7%;background:linear-gradient(135deg,#dff3eb,#fff4dc)}
    h1{font-size:44px;margin:0 0 16px;max-width:760px}.lead{font-size:18px;line-height:1.8;max-width:720px}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(240px,1fr));gap:18px;padding:36px 7% 70px}
    .card{background:#fff;padding:24px;border-radius:18px;border:1px solid var(--line);box-shadow:0 8px 30px #28352b12}.price{font-size:26px;font-weight:800;color:var(--green)}
    .meta{font-size:13px;color:#64736d}.button{display:inline-block;margin-top:12px;padding:10px 16px;border-radius:999px;background:var(--green);color:white;text-decoration:none}footer{padding:24px 7%;background:#17221f;color:#dce7e3}
    """
    write(static / "styles.css", static_css)
    write(static / "index.html", """
    <!doctype html><html lang="ja"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>青葉ダイニング｜季節のイベント</title><link rel="stylesheet" href="styles.css"></head>
    <body data-page-type="static-html"><header><div class="brand">青葉ダイニング</div><div class="tag">TRAINING STATIC SITE</div></header>
    <main><section class="hero"><p class="meta">更新日：2026-09-01</p><h1>秋の味覚フェア 2026</h1><p class="lead">国産きのこと旬魚を使った期間限定コース。平日夜は乾杯ドリンク付きです。</p></section>
    <section class="grid" id="campaign-list">
      <article class="card" data-id="AOBA-CP-001"><p class="meta">2026-09-01〜2026-10-15</p><h2>秋季限定コース</h2><p class="price" data-price="3980">3,980円</p><p>前菜・旬魚・きのこ御飯・甘味の全6品。</p><a class="button" href="menu.html#autumn">詳細を見る</a></article>
      <article class="card" data-id="AOBA-CP-002"><p class="meta">毎週 月〜木</p><h2>平日ペアディナー</h2><p class="price" data-price="6800">2名 6,800円</p><p>17時までの予約限定。祝前日は対象外。</p><a class="button" href="menu.html#pair">詳細を見る</a></article>
      <article class="card" data-id="AOBA-CP-003"><p class="meta">2026-09-10〜2026-09-30</p><h2>敬老の日 お持ち帰り膳</h2><p class="price" data-price="2480">2,480円</p><p>受取日の2日前までに要予約。</p><a class="button" href="menu.html#takeout">詳細を見る</a></article>
    </section></main><footer>本ページは研修用の合成サイトです。実在の店舗とは関係ありません。</footer></body></html>
    """)
    write(static / "menu.html", """
    <!doctype html><html lang="ja"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>青葉ダイニング｜メニュー詳細</title><link rel="stylesheet" href="styles.css"></head>
    <body data-page-type="static-html"><header><div class="brand">青葉ダイニング</div><a href="index.html">イベント一覧へ</a></header><main class="grid">
    <article class="card" id="autumn"><h2>秋季限定コース</h2><p class="price">3,980円</p><p>提供期間：2026-09-01〜2026-10-15</p><p>条件：2名様より。前日18時までの予約。</p></article>
    <article class="card" id="pair"><h2>平日ペアディナー</h2><p class="price">2名 6,800円</p><p>提供日：月〜木（祝前日を除く）</p><p>条件：17時までの予約限定。</p></article>
    <article class="card" id="takeout"><h2>敬老の日 お持ち帰り膳</h2><p class="price">2,480円</p><p>提供期間：2026-09-10〜2026-09-30</p><p>条件：受取日の2日前までに予約。</p></article></main></body></html>
    """)

    activities = [
        {"id":"KAZE-CP-101","title":"月見テラスプラン","price_yen":4500,"start_date":"2026-09-05","end_date":"2026-09-29","status":"受付中","conditions":"18:00以降・2名以上","updated_at":"2026-09-02T18:30:00+09:00"},
        {"id":"KAZE-CP-102","title":"週末ファミリーランチ","price_yen":2200,"start_date":"2026-09-06","end_date":"2026-10-25","status":"受付中","conditions":"土日祝・小学生以下ドリンク無料","updated_at":"2026-09-02T18:30:00+09:00"},
        {"id":"KAZE-CP-103","title":"早割 忘年会予約","price_yen":5000,"start_date":"2026-09-01","end_date":"2026-09-30","status":"受付中","conditions":"11月開催・8名以上・9月中予約","updated_at":"2026-09-03T07:15:00+09:00"},
        {"id":"KAZE-CP-104","title":"平日14時のデザートセット","price_yen":1380,"start_date":"2026-09-01","end_date":"2026-11-30","status":"受付中","conditions":"平日14:00〜16:30","updated_at":"2026-09-03T07:15:00+09:00"},
        {"id":"KAZE-CP-105","title":"料理長おまかせ七皿","price_yen":7200,"start_date":"2026-09-12","end_date":"2026-10-31","status":"残席わずか","conditions":"3日前までに予約","updated_at":"2026-09-03T08:00:00+09:00"},
        {"id":"KAZE-CP-106","title":"秋のテイクアウトBOX","price_yen":1980,"start_date":"2026-09-15","end_date":"2026-10-15","status":"準備中","conditions":"店頭受取のみ","updated_at":"2026-09-03T08:00:00+09:00"},
        {"id":"KAZE-CP-107","title":"夏のビアテラス","price_yen":4800,"start_date":"2026-07-01","end_date":"2026-08-31","status":"終了","conditions":"雨天中止","updated_at":"2026-09-01T09:00:00+09:00"},
    ]
    write(dynamic / "activities.json", json.dumps({"generated_at":CAPTURE_BASE,"items":activities}, ensure_ascii=False, indent=2))
    write(dynamic / "styles.css", static_css + "\n.loader{text-align:center;padding:30px}.hidden{display:none}.status{font-weight:700;color:#a14c14}")
    write(dynamic / "index.html", """
    <!doctype html><html lang="ja"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>風音キッチン｜イベント</title><link rel="stylesheet" href="styles.css"></head>
    <body data-page-type="javascript-rendered"><header><div class="brand">風音キッチン</div><div class="tag">TRAINING DYNAMIC SITE</div></header>
    <main><section class="hero"><p class="meta">この一覧は JavaScript 実行後に表示されます</p><h1>EVENTS & OFFERS</h1><p class="lead">最初のHTMLには活動カードがありません。JSON取得後、3件ずつ表示します。</p></section>
    <div id="loading" class="loader">データを読み込み中...</div><section id="activity-list" class="grid" aria-live="polite"></section><div class="loader"><button id="more" class="button hidden">もっと見る</button></div></main>
    <footer>研修用合成サイト。ログインや検証回避は不要です。</footer><script src="app.js"></script></body></html>
    """)
    write(dynamic / "app.js", """
    const list=document.querySelector('#activity-list'),loading=document.querySelector('#loading'),more=document.querySelector('#more');
    let items=[],shown=0; const batch=3;
    const esc=s=>String(s).replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
    function render(){const next=items.slice(shown,shown+batch);next.forEach(x=>{const el=document.createElement('article');el.className='card activity-card';el.dataset.id=x.id;el.innerHTML=`<p class="meta">${esc(x.start_date)}〜${esc(x.end_date)}</p><h2>${esc(x.title)}</h2><p class="price" data-price="${x.price_yen}">${x.price_yen.toLocaleString('ja-JP')}円</p><p class="status">${esc(x.status)}</p><p>${esc(x.conditions)}</p><p class="meta">更新：${esc(x.updated_at)}</p>`;list.appendChild(el)});shown+=next.length;more.classList.toggle('hidden',shown>=items.length)}
    fetch('activities.json').then(r=>{if(!r.ok)throw new Error(`HTTP ${r.status}`);return r.json()}).then(d=>{items=d.items;loading.remove();render()}).catch(e=>{loading.textContent=`読み込み失敗：${e.message}。HTTPサーバー経由で開いてください。`});
    more.addEventListener('click',render);
    """)

    reviews = [
        {"review_id":"REV-9001","rating":5,"review_date":"2026-09-01","author_alias":"M.K.","text":"季節の野菜がきれいで、店員さんの説明も丁寧でした。記念日プレートも嬉しかったです。","owner_reply_status":"未返信","risk":"LOW"},
        {"review_id":"REV-9002","rating":2,"review_date":"2026-09-01","author_alias":"Sato77","text":"金曜19時、予約していたのに料理が出るまで35分待ちました。味は良かったですが説明がありませんでした。","owner_reply_status":"未返信","risk":"MEDIUM"},
        {"review_id":"REV-9003","rating":4,"review_date":"2026-08-30","author_alias":"旅する猫","text":"ランチはコスパが良いです。入口が少し分かりにくいので案内板があると助かります。","owner_reply_status":"返信済み","risk":"LOW"},
        {"review_id":"REV-9004","rating":1,"review_date":"2026-08-29","author_alias":"匿名希望","text":"会計がメニュー表示より高く感じました。明細を受け取っていないので確認してほしいです。","owner_reply_status":"要確認","risk":"HIGH"},
        {"review_id":"REV-9005","rating":5,"review_date":"2026-08-27","author_alias":"Yumi","text":"子どもへの対応が優しく、アレルギーについても厨房に確認してくれました。","owner_reply_status":"未返信","risk":"LOW"},
        {"review_id":"REV-9006","rating":3,"review_date":"2026-08-26","author_alias":"K.T.","text":"料理は美味しいですが、QRメニューの文字が小さく注文しづらかったです。","owner_reply_status":"未返信","risk":"MEDIUM"},
        {"review_id":"REV-9007","rating":2,"review_date":"2026-08-24","author_alias":"Rin","text":"テイクアウトの受取時間を20分過ぎました。混雑時の案内を改善してほしいです。","owner_reply_status":"未返信","risk":"MEDIUM"},
        {"review_id":"REV-9008","rating":4,"review_date":"2026-08-23","author_alias":"Nori","text":"魚料理が特に良かったです。次は季節コースを予約したいです。","owner_reply_status":"返信済み","risk":"LOW"},
    ]
    write(review / "reviews.json", json.dumps({"platform":"MapReview Training","place_id":"TRAINING-HIKARI-001","generated_at":CAPTURE_BASE,"authorized_training_export":True,"items":reviews}, ensure_ascii=False, indent=2))
    with (review / "reviews_export.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(reviews[0]))
        w.writeheader(); w.writerows(reviews)
    write(review / "reviews.html", """
    <!doctype html><html lang="ja"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>MapReview Training｜光庭ビストロ</title><link rel="stylesheet" href="../02_静态竞品官网/styles.css"></head>
    <body data-page-type="authorized-platform-simulation"><header><div class="brand">MapReview Training</div><div class="tag">AUTHORIZED TRAINING VIEW</div></header><main>
    <section class="hero"><p class="meta">研修用・実在プラットフォームではありません</p><h1>光庭ビストロ</h1><p class="lead">表示中のレビューだけを読み取り、ログイン・验证码・アクセス制限が現れた場合は停止してください。</p></section>
    <section id="review-list" class="grid"></section><div style="text-align:center;padding:0 0 48px"><button id="more" class="button">さらに表示</button></div></main>
    <footer>TRAINING DATA ONLY｜公開投稿・返信操作は実行しません。</footer><script src="reviews.js"></script></body></html>
    """)
    write(review / "reviews.js", """
    const list=document.querySelector('#review-list'),more=document.querySelector('#more');let items=[],shown=0;
    function render(){items.slice(shown,shown+4).forEach(x=>{const e=document.createElement('article');e.className='card review-card';e.dataset.reviewId=x.review_id;e.innerHTML=`<p class="meta">${x.review_date}｜${x.author_alias}</p><h2>${'★'.repeat(x.rating)}${'☆'.repeat(5-x.rating)}</h2><p class="review-text">${x.text}</p><p class="meta">返信状態：${x.owner_reply_status}</p>`;list.appendChild(e)});shown+=4;if(shown>=items.length)more.hidden=true}
    fetch('reviews.json').then(r=>r.json()).then(d=>{items=d.items;render()}).catch(e=>{list.innerHTML=`<p>読み込み失敗：${e.message}</p>`});more.addEventListener('click',render);
    """)
    write(review / "README_授权说明.md", """
    # 评论平台训练页授权说明

    - 本目录模拟 Google 地图等评论平台的“已授权可见页面”，不是任何真实平台的镜像。
    - `reviews_export.csv` 与 `reviews.json` 代表店铺后台或官方接口提供的训练导出。
    - 课堂优先顺序：官方 API/后台导出 → 已授权浏览器读取当前可见内容 → 人工导出。
    - 禁止绕过登录、验证码、反爬限制或权限范围；出现限制时记录证据并停止。
    - 评论者仅使用别名，不采集私人联系方式或个人主页信息。
    """)


def build_docs(root: Path) -> None:
    write(root / "README.md", """
    # S-S07-C01 多来源竞品监控｜原始资料

    本资料包用于《Codex销售支援16小时实战训练营》第7课主案例。数据均为合成数据，适合课堂演示和学员练习。

    ## 文件用途

    1. `01_竞品URL清单.xlsx`：入口、预期路线、权限和停止条件。
    2. `02_静态竞品官网/`：HTML 源代码直接含活动数据，适合 HTTP GET + DOM/CSS。
    3. `03_JS动态活动页/`：初始 HTML 无活动卡片，运行 JavaScript 后从 JSON 加载，并需点击“更多”。
    4. `04_评论平台训练页/`：已授权平台页面模拟、JSON 与 CSV 导出，不允许绕过权限。
    5. `05_PDF竞品菜单/`：同一菜单的文本型与扫描型变体，用于判断直接提取或 OCR。
    6. `06_竞品监控字段模板.xlsx`：诊断、统一监控、活动和评论的空白模板及字段字典。
    7. `07_公司资料包.md`：监控目标、竞品范围和业务口径。
    8. `08_网页访问与停止规则.md`：合规边界、失败处理与证据要求。
    9. `09_字段说明与练习要求.md`：课堂任务、字段映射和验收标准。

    ## 在线查看训练网页

    网站发布后，可直接打开 [静态竞品官网](02_静态竞品官网/index.html)、[JS 动态活动页](03_JS动态活动页/index.html)、[评论平台训练页](04_评论平台训练页/reviews.html)，无需启动本地服务器。课堂抓取练习请使用资料库提供的原始链接。

    ## 本地练习（可选）

    在本目录执行：

    ```bash
    python3 -m http.server 8000
    ```

    本地练习时，打开 `http://127.0.0.1:8000/02_静态竞品官网/index.html` 等相应目录地址；Excel 清单默认提供线上地址。不要直接双击动态页；`file://` 下浏览器通常会拦截 JSON 请求，这正好也可用于讲解“页面能打开但数据未加载”的差异。

    ## 重新生成整套资料

    生成脚本默认会覆盖本目录中同名的训练文件。请先复制出自己的课堂修改版，再执行：

    ```bash
    python3 -m pip install -r requirements.txt
    # macOS：brew install poppler
    # Ubuntu/Debian：sudo apt-get install poppler-utils
    python3 S-S07-C01_生成原始资料.py
    ```

    如果只需要重新生成网页、文档与 PDF，不改动 Excel，可加 `--skip-xlsx`。生成本地入口时加 `--base-url http://127.0.0.1:8000`。

    ## 预期课堂产出（不随原始资料预填）

    - 网页类型诊断表.xlsx
    - 抓取路线图.md
    - 竞品活动与评论表.xlsx
    - 首轮快照 JSON/CSV
    """)
    write(root / "07_公司资料包.md", """
    # 株式会社みなと食研｜竞品监控公司资料包

    ## 公司背景

    - 自有品牌：港灯餐厅（みなと灯りダイニング）
    - 业态：东京近郊的家庭型西餐与季节套餐
    - 核心客群：25–55岁家庭、附近企业聚餐、周末游客
    - 当前目标：观察秋季活动、套餐价格、预约条件和顾客对等待时间/服务的反馈

    ## 本次监控对象（均为合成企业）

    | 竞品ID | 名称 | 主要来源 | 监控重点 |
    |---|---|---|---|
    | CMP-A | 青葉ダイニング | 静态官网 | 活动名、价格、日期、预约条件 |
    | CMP-B | 風音キッチン | JS动态页 | 分批加载的活动、状态、更新时间 |
    | CMP-C | 光庭ビストロ | 评论平台训练页 | 评分、日期、原文、风险线索 |
    | CMP-D | 光庭ビストロ | PDF菜单 | 菜品、价格、适用时段、脚注条件 |

    ## 业务口径

    - 金额统一存为整数日元，显示可使用 `#,##0円`。
    - 日期存为 `YYYY-MM-DD`；区间拆成开始日与结束日。
    - “每周/长期/未公布”等无法转成确定日期的信息保留原文，并在备注中标记“不确定”。
    - 竞品推测与页面事实分开记录；不得把评论者意见当作已证实事实。
    - 首轮快照只建立基线，不直接给出经营结论。

    ## 权限与责任

    - 课堂账号仅允许读取本地训练页面和提供的导出文件。
    - 不登录真实平台、不公开回复、不采集私人信息、不突破验证码或访问限制。
    - 价格解释、竞争判断、对外发布和 CRM 正式写入均需销售负责人确认。
    """)
    write(root / "08_网页访问与停止规则.md", """
    # 网页访问与停止规则

    ## 允许动作

    - 读取本地训练网页、JSON、CSV 与 PDF。
    - 对静态页发起有限次数 HTTP GET，并解析明确字段。
    - 在本地动态页等待元素、点击“更多”、读取已渲染 DOM、保存截图。
    - 读取已提供的评论导出与当前可见训练评论。
    - 下载/复制训练 PDF，判断是否存在文本层，再决定直接提取或 OCR。

    ## 必须停止并记录

    | 触发条件 | 立即动作 | 记录内容 | 后续负责人 |
    |---|---|---|---|
    | 出现真实登录页或账号要求 | 停止自动化 | URL、时间、截图、最后一步 | 账号管理员 |
    | 出现验证码/反爬提示 | 不重试绕过 | 页面提示、请求次数、截图 | 讲师/合规负责人 |
    | 访问被拒绝（401/403） | 停止请求 | 状态码、URL、时间 | 数据负责人 |
    | 页面条款禁止自动采集 | 改用 API/导出/人工 | 条款位置、替代路线 | 合规负责人 |
    | 动态选择器失效 | 停止写入，不覆盖旧快照 | 错误日志、失败截图 | 脚本维护者 |
    | PDF 无文本且 OCR 低置信 | 标记失败/待确认 | 页码、截图、疑似字段 | 业务复核人 |
    | 来源与字段冲突 | 不自行判断真伪 | 两处原文、URL/页码 | 销售负责人 |

    ## 请求控制

    - 只访问清单内训练入口；最大翻页/点击次数为 10。
    - 每个入口失败最多重试 2 次；重复失败后进入人工接管。
    - 写入新结果前先保留旧快照；失败时不得用空表覆盖旧数据。
    - 每条结构化记录都必须包含 `source_url`、`captured_at`、`evidence_text`。
    """)
    write(root / "09_字段说明与练习要求.md", """
    # 字段说明与练习要求

    ## 练习目标

    先诊断来源，再选 HTTP、浏览器、平台/API 或 PDF/OCR 路线，最后统一写入一张竞品监控表，并保存首轮快照。

    ## 必做步骤

    1. 从案例资料页打开训练网页，逐个检查 `01_竞品URL清单.xlsx` 中的入口；本地练习时也可启动 HTTP 服务器。
    2. 只观察页面源代码、渲染状态、滚动/点击、登录状态和文件类型，先填写“网页诊断模板”。
    3. 为每个来源写出首选路线、备选路线、停止条件和证据保存方式。
    4. 分别提取活动、评论和菜单字段；不得用同一脚本硬套全部来源。
    5. 标准化日期、金额、来源类型和状态，保留原文证据。
    6. 输出 `snapshot_2026-09-03.json` 或 CSV，作为后续差分基线。

    ## 统一字段（核心）

    | 字段 | 类型 | 说明 |
    |---|---|---|
    | record_id | 文本 | 自建唯一键，如来源ID+对象ID |
    | competitor_id | 文本 | CMP-A～CMP-D |
    | source_type | 枚举 | STATIC / DYNAMIC / PLATFORM / PDF |
    | item_type | 枚举 | ACTIVITY / REVIEW / MENU |
    | title | 文本 | 活动、评论主题或菜品名称 |
    | price_yen | 整数 | 日元金额，无则留空 |
    | start_date / end_date | 日期 | 无法确定时留空并保留原文 |
    | rating | 数值 | 1–5，仅评论使用 |
    | status | 文本 | 受付中、終了、未返信等原始状态 |
    | source_url | 文本 | URL、文件路径或导出来源 |
    | captured_at | 日期时间 | 带时区的抓取时间 |
    | evidence_text | 文本 | 支持本行字段的短原文 |
    | evidence_locator | 文本 | CSS选择器、review_id、PDF页码等 |
    | result_status | 枚举 | SUCCESS / PARTIAL / FAILED / STOPPED |
    | failure_reason | 文本 | 失败或停止时必填 |

    ## 验收

    - 四类来源的技术路线选择正确。
    - 每条结果均可回到 URL/文件、抓取时间和原文证据。
    - 动态页完成等待与“更多”加载，结果按业务主键去重。
    - 评论页优先使用提供的授权导出，不采集私人资料。
    - 文本 PDF 直接提取；扫描 PDF 说明 OCR 与页码证据。
    - 失败能够说明原因，且旧快照没有被覆盖。
    """)
    write(root / "启动训练网页.sh", """
    #!/usr/bin/env bash
    set -euo pipefail
    cd "$(dirname "$0")"
    echo "Open: http://127.0.0.1:8000/"
    python3 -m http.server 8000 --bind 127.0.0.1
    """)
    (root / "启动训练网页.sh").chmod(0o755)
    write(root / "requirements.txt", """
    openpyxl>=3.1,<4
    reportlab>=4,<5
    """)


def pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    regular = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    bold = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    pdfmetrics.registerFont(TTFont("MenuSans", regular))
    pdfmetrics.registerFont(TTFont("MenuSansBold", bold))
    return "MenuSans", "MenuSansBold"


def draw_menu_pdf(path: Path) -> None:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    regular, bold = pdf_fonts(); path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=A4); w, h = A4
    c.setFillColor(HexColor("#173F35")); c.rect(0, h-115, w, 115, stroke=0, fill=1)
    c.setFillColor(HexColor("#FFFFFF")); c.setFont(bold, 23); c.drawString(48, h-62, "HIKARI BISTRO")
    c.setFont(regular, 11); c.drawString(48, h-88, "2026 AKI NO MENU / AUTUMN MENU")
    items = [
        ("Kisetsu Yasai Appetizer", "JPY 980", "Lunch & Dinner"),
        ("Aki-zake Herb Roast", "JPY 2,480", "Dinner only"),
        ("Kinoko Cream Pasta", "JPY 1,780", "All day"),
        ("Beef Cheek Red Wine Stew", "JPY 2,980", "Limited quantity"),
        ("Kuri & Hojicha Parfait", "JPY 1,280", "From 14:00"),
        ("Autumn Share Course", "JPY 5,600", "2+ guests; reserve one day ahead"),
    ]
    y=h-155
    for i,(name,price,note) in enumerate(items,1):
        c.setFillColor(HexColor("#173F35")); c.setFont(bold,13); c.drawString(52,y,name)
        c.setFont(bold,13); c.drawRightString(w-52,y,price)
        c.setFillColor(HexColor("#66756F")); c.setFont(regular,9); c.drawString(52,y-17,note)
        c.setStrokeColor(HexColor("#D8E1DD")); c.line(52,y-27,w-52,y-27); y-=62
    c.setFillColor(HexColor("#7C4A23")); c.setFont(regular,9)
    c.drawString(52,82,"Prices include tax. Items may change depending on availability.")
    c.drawString(52,66,"Please confirm allergy support with the restaurant when booking.")
    c.setFillColor(HexColor("#777777")); c.setFont(regular,8); c.drawString(52,40,"SYNTHETIC TRAINING MATERIAL | Page 1/1 | Updated 2026-09-02")
    c.save()


def make_scanned_pdf(text_pdf: Path, scan_pdf: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    exe = shutil.which("pdftoppm")
    if not exe:
        raise RuntimeError("未找到 pdftoppm；请安装 Poppler 后生成扫描型 PDF")
    with tempfile.TemporaryDirectory() as td:
        prefix = Path(td) / "page"
        subprocess.run([exe,"-png","-r","150",str(text_pdf),str(prefix)],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
        png = sorted(Path(td).glob("page-*.png"))[0]
        c=canvas.Canvas(str(scan_pdf),pagesize=A4); w,h=A4
        c.drawImage(str(png),0,0,width=w,height=h,mask='auto'); c.save()


def build_pdfs(root: Path) -> None:
    folder=root/"05_PDF竞品菜单"; folder.mkdir(parents=True,exist_ok=True)
    text_pdf=folder/"光庭_菜单_文本型.pdf"; scan_pdf=folder/"光庭_菜单_扫描型.pdf"
    draw_menu_pdf(text_pdf); make_scanned_pdf(text_pdf,scan_pdf)


def style_sheet(ws, widths):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; ws.sheet_view.showGridLines=False
    for cell in ws[1]:
        cell.fill=PatternFill("solid",fgColor="173F35"); cell.font=Font(color="FFFFFF",bold=True); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws.row_dimensions[1].height=34
    for col,width in widths.items(): ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)


def build_xlsx(root: Path, base_url: str = DEFAULT_BASE_URL) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    wb=Workbook(); ws=wb.active; ws.title="URL清单"
    headers=["source_id","competitor_id","竞品名称","来源名称","source_type","训练入口","是否需本地服务器","登录/授权","首选路线","备选路线","预期字段","证据要求","停止条件","备注"]
    ws.append(headers)
    rows=[
      ["SRC-A-01","CMP-A","青葉ダイニング","静态活动官网","STATIC",f"{base_url}/02_静态竞品官网/index.html","否（线上访问）" if not base_url.startswith("http://127.0.0.1") else "是","无需","HTTP GET + DOM/CSS","浏览器查看源代码","活动/价格/日期/条件","URL+captured_at+CSS选择器+原文","401/403/登录/条款限制","正文存在于初始HTML"],
      ["SRC-B-01","CMP-B","風音キッチン","JS活动列表","DYNAMIC",f"{base_url}/03_JS动态活动页/index.html","否（线上访问）" if not base_url.startswith("http://127.0.0.1") else "是","无需","Playwright/Chrome 等待+点击","检查JSON/XHR是否可用","活动/价格/日期/状态/更新时间","URL+captured_at+截图+data-id","验证码/登录/选择器连续失败","初始HTML没有活动卡片"],
      ["SRC-C-01","CMP-C","光庭ビストロ","评论平台训练页","PLATFORM",f"{base_url}/04_评论平台训练页/reviews.html","否（线上访问）" if not base_url.startswith("http://127.0.0.1") else "是","已授权训练视图","官方API/后台导出","授权浏览器读取可见评论","评分/日期/别名/正文/回复状态","平台名+review_id+captured_at+原文","真实登录/验证码/反爬/无权限","优先读取CSV或JSON导出"],
      ["SRC-D-01","CMP-D","光庭ビストロ","文本型PDF菜单","PDF",f"{base_url}/05_PDF竞品菜单/光庭_菜单_文本型.pdf","否","无需","下载+文本层提取","逐页人工核对","菜名/价格/时段/条件","文件名+页码+captured_at+原文","文件损坏/来源不明","应可直接复制文本"],
      ["SRC-D-02","CMP-D","光庭ビストロ","扫描型PDF菜单","PDF",f"{base_url}/05_PDF竞品菜单/光庭_菜单_扫描型.pdf","否","无需","逐页OCR","人工录入并双人复核","菜名/价格/时段/条件","文件名+页码+OCR片段+置信说明","OCR低置信且无法人工确认","与文本版内容相同，仅用于路线判断"],
    ]
    for r in rows: ws.append(r)
    style_sheet(ws,{"A":14,"B":14,"C":20,"D":22,"E":14,"F":58,"G":18,"H":18,"I":28,"J":24,"K":30,"L":34,"M":32,"N":34})
    wb.save(root/"01_竞品URL清单.xlsx")

    wb=Workbook(); diag=wb.active; diag.title="网页诊断模板"
    diag.append(["source_id","入口/文件","初始HTML有目标正文","需JS渲染","需滚动/点击","登录/平台授权","文件类型","判断source_type","首选路线","备选路线","停止条件","诊断证据","诊断人","诊断时间"])
    for _ in range(8): diag.append([""]*14)
    style_sheet(diag,{"A":14,"B":45,"C":20,"D":16,"E":18,"F":20,"G":14,"H":20,"I":28,"J":24,"K":32,"L":32,"M":14,"N":22})
    for col in ("C","D","E"): diag.add_data_validation(DataValidation(type="list",formula1='"是,否,不确定"',allow_blank=True)); diag.data_validations.dataValidation[-1].add(f"{col}2:{col}200")
    diag.add_data_validation(DataValidation(type="list",formula1='"STATIC,DYNAMIC,PLATFORM,PDF,STOP"',allow_blank=True)); diag.data_validations.dataValidation[-1].add("H2:H200")

    unified=wb.create_sheet("统一监控表模板")
    unified.append(["record_id","competitor_id","competitor_name","source_id","source_type","item_type","item_key","title","price_yen","start_date","end_date","rating","status","conditions_or_summary","source_url","captured_at","evidence_text","evidence_locator","result_status","failure_reason","reviewer","confirmed_at"])
    for _ in range(12): unified.append([""]*22)
    style_sheet(unified,{"A":19,"B":15,"C":20,"D":15,"E":14,"F":14,"G":20,"H":28,"I":14,"J":14,"K":14,"L":10,"M":16,"N":32,"O":46,"P":24,"Q":48,"R":28,"S":16,"T":36,"U":15,"V":22})
    unified.add_data_validation(DataValidation(type="list",formula1='"STATIC,DYNAMIC,PLATFORM,PDF"',allow_blank=True)); unified.data_validations.dataValidation[-1].add("E2:E500")
    unified.add_data_validation(DataValidation(type="list",formula1='"ACTIVITY,REVIEW,MENU"',allow_blank=True)); unified.data_validations.dataValidation[-1].add("F2:F500")
    unified.add_data_validation(DataValidation(type="list",formula1='"SUCCESS,PARTIAL,FAILED,STOPPED"',allow_blank=True)); unified.data_validations.dataValidation[-1].add("S2:S500")
    for cell in unified[1]:
        alignment = copy(cell.alignment)
        alignment.text_rotation = 45
        cell.alignment = alignment

    fields=wb.create_sheet("字段字典")
    fields.append(["字段","必填条件","数据类型","允许值/格式","说明","示例"])
    dictionary=[
      ["record_id","全部","文本","唯一","统一记录主键","SRC-A-01|AOBA-CP-001"],["competitor_id","全部","文本","CMP-A～D","竞品编号","CMP-A"],
      ["source_type","全部","枚举","STATIC/DYNAMIC/PLATFORM/PDF","来源类型","STATIC"],["item_type","全部","枚举","ACTIVITY/REVIEW/MENU","对象类型","ACTIVITY"],
      ["price_yen","活动/菜单可选","整数",">=0","纯数字日元","3980"],["start_date","活动可选","日期","yyyy-mm-dd","未知则留空","2026-09-01"],
      ["rating","评论必填","数值","1～5","评论评分","2"],["source_url","全部","文本","URL或相对文件路径","可回到原始来源",f"{base_url}/02_静态竞品官网/index.html"],
      ["captured_at","全部","日期时间","ISO 8601含时区","抓取/提取时间",CAPTURE_BASE],["evidence_text","全部","文本","短原文","直接支持本行字段","秋季限定コース 3,980円"],
      ["evidence_locator","全部","文本","CSS/review_id/页码","证据定位","article[data-id=AOBA-CP-001]"],["result_status","全部","枚举","SUCCESS/PARTIAL/FAILED/STOPPED","执行结果","SUCCESS"],
      ["failure_reason","失败/停止","文本","自由文本","不得只写失败","HTTP 403，已停止请求"],
    ]
    for r in dictionary: fields.append(r)
    style_sheet(fields,{"A":22,"B":18,"C":14,"D":30,"E":38,"F":42})
    wb.save(root/"06_竞品监控字段模板.xlsx")


def main() -> None:
    ap=argparse.ArgumentParser(); ap.add_argument("--output",type=Path,default=Path(__file__).resolve().parent); ap.add_argument("--skip-xlsx",action="store_true"); ap.add_argument("--base-url",default=DEFAULT_BASE_URL); args=ap.parse_args()
    root=args.output.resolve(); root.mkdir(parents=True,exist_ok=True)
    build_pages(root); build_docs(root); build_pdfs(root)
    if not args.skip_xlsx: build_xlsx(root, args.base_url.rstrip("/"))
    print(f"Generated {CASE_ID} materials at: {root}")


if __name__ == "__main__": main()
